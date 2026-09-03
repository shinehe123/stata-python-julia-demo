"""Transform PSS bibliographic records into patent-company and edge tables."""

from __future__ import annotations

import csv
import json
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Iterator

from .normalize import normalize_name, split_applicants

ALIASES = {
    "application_number": ("申请号", "申请（专利）号", "application_number", "申请号码"),
    "publication_number": ("公开号", "公开（公告）号", "publication_number", "公开公告号"),
    "title": ("发明名称", "名称", "title", "专利名称"),
    "application_date": ("申请日", "application_date"),
    "publication_date": ("公开（公告）日", "公开日", "publication_date"),
    "applicants": ("申请（专利权）人", "申请人", "专利权人", "applicants"),
    "ipc": ("IPC分类号", "IPC主分类号", "ipc", "分类号"),
    "patent_type": ("专利类型", "申请类型", "patent_type"),
}


@dataclass(frozen=True)
class Company:
    company_id: str
    name: str
    aliases: tuple[str, ...]

    @property
    def match_names(self) -> set[str]:
        return {normalize_name(x) for x in (self.name, *self.aliases) if x}


def load_companies(path: Path) -> list[Company]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        rows = csv.DictReader(handle)
        required = {"company_id", "company_name"}
        if not rows.fieldnames or not required.issubset(rows.fieldnames):
            raise ValueError("公司表必须包含 company_id 和 company_name 列")
        return [
            Company(
                str(row["company_id"]).strip(),
                str(row["company_name"]).strip(),
                tuple(split_applicants(row.get("aliases", ""))),
            )
            for row in rows
            if str(row.get("company_id", "")).strip()
        ]


def iter_records(path: Path) -> Iterator[dict[str, str]]:
    suffix = path.suffix.lower()
    if suffix in {".csv", ".txt"}:
        with path.open(encoding="utf-8-sig", newline="") as handle:
            yield from csv.DictReader(handle)
        return
    if suffix in {".json", ".jsonl"}:
        with path.open(encoding="utf-8-sig") as handle:
            if suffix == ".jsonl":
                for line in handle:
                    if line.strip():
                        yield json.loads(line)
            else:
                data = json.load(handle)
                yield from (data if isinstance(data, list) else data.get("records", []))
        return
    if suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("读取 xlsx 需要安装可选依赖：pip install '.[excel]'") from exc
        sheet = load_workbook(path, read_only=True, data_only=True).active
        rows = sheet.iter_rows(values_only=True)
        headers = [str(v or "").strip() for v in next(rows)]
        for values in rows:
            yield {header: str(value or "") for header, value in zip(headers, values)}
        return
    raise ValueError(f"不支持的文件类型：{suffix}")


def canonicalize(row: dict[str, object]) -> dict[str, str]:
    clean = {str(k).strip(): str(v or "").strip() for k, v in row.items()}
    return {
        field: next((clean[name] for name in names if clean.get(name)), "")
        for field, names in ALIASES.items()
    }


def build_tables(
    records: Iterable[dict[str, object]], companies: list[Company]
) -> tuple[list[dict[str, str]], list[dict[str, object]]]:
    """Return patent-level memberships and aggregated undirected edges."""
    company_by_name = {
        alias: company for company in companies for alias in company.match_names
    }
    patents: dict[str, dict[str, str]] = {}
    edges: defaultdict[tuple[str, str, str], set[str]] = defaultdict(set)

    for raw in records:
        record = canonicalize(raw)
        patent_id = record["application_number"] or record["publication_number"]
        if not patent_id:
            continue
        applicants = split_applicants(record["applicants"])
        listed = {
            company_by_name[normalize_name(name)]
            for name in applicants
            if normalize_name(name) in company_by_name
        }
        if not listed:
            continue
        year = (record["application_date"] or record["publication_date"])[:4]
        for company in sorted(listed, key=lambda item: item.company_id):
            key = f"{patent_id}\0{company.company_id}"
            patents[key] = {
                **record,
                "patent_id": patent_id,
                "company_id": company.company_id,
                "company_name": company.name,
                "applicant_count": str(len(applicants)),
                "is_collaboration": str(len(applicants) >= 2).lower(),
            }
            for partner in applicants:
                if normalize_name(partner) not in company.match_names:
                    edges[(company.company_id, partner, year)].add(patent_id)

    edge_rows = [
        {
            "company_id": company_id,
            "partner_name": partner,
            "year": year,
            "joint_patent_count": len(patent_ids),
        }
        for (company_id, partner, year), patent_ids in sorted(edges.items())
    ]
    return list(patents.values()), edge_rows


def write_csv(path: Path, rows: list[dict[str, object]], fieldnames: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)

